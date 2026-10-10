import copy
import io
import unittest

import openpyxl
import pandas as pd

import db
import execution as ex
import reports
import versions as vs
from tests import fixtures as fx


class DivisionReportTest(unittest.TestCase):
    def test_summary_and_excel_tie_to_total(self):
        cur = pd.DataFrame(fx.make_lines())
        prev = pd.DataFrame(fx.make_lines(fx.SPEC[:3]))
        s = reports.summary_by_division(cur, prev, fx.master())
        self.assertAlmostEqual(s["Tổng"].sum(), cur["total_budget"].sum())
        self.assertAlmostEqual(s["Năm trước"].sum(), prev["total_budget"].sum())
        parts = s[["TSCĐ (CAPEX)", "CCDC", "OPEX", "Chưa phân loại"]].sum(axis=1)
        self.assertTrue(((parts - s["Tổng"]).abs() < 1).all())
        self.assertAlmostEqual(s["Tỷ trọng"].sum(), 1.0)
        wb = openpyxl.load_workbook(io.BytesIO(reports.report_xlsx(cur, prev, fx.master(), {"budget_year": fx.YEAR, "scope": "thử"})))
        ws = wb["Khối - Phòng - Hạng mục"]
        items = [r for r in ws.iter_rows(min_row=5) if ws.row_dimensions[r[0].row].outline_level == 2]
        self.assertAlmostEqual(sum(r[7].value for r in items), cur["total_budget"].sum())
        self.assertEqual(ws.cell(row=ws.max_row, column=1).value, "TỔNG CỘNG")

    def test_empty(self):
        self.assertTrue(reports.summary_by_division(pd.DataFrame(), None, fx.master()).empty)
        self.assertTrue(reports.report_xlsx(pd.DataFrame(), None, fx.master(), {"budget_year": fx.YEAR}))


class ExecutionTest(unittest.TestCase):
    def setUp(self):
        self.lines = pd.DataFrame(fx.make_lines()).assign(site_code=fx.SITE)
        self.lap = self.lines[(self.lines["dept_proposing"] == fx.KD) & (self.lines["catalog_code"] == fx.LAPTOP)].iloc[0]
        self.zw = self.lines[self.lines["catalog_code"] == fx.ZWCAD].iloc[0]
        b = self.lap["total_budget"]
        self.E = [
            {"id": 1, "site_code": fx.SITE, "item_code": self.lap["item_code"], "kind": db.EXEC_REQUEST, "amount": b, "doc_date": "2026-10-05"},
            {"id": 2, "site_code": fx.SITE, "item_code": self.lap["item_code"], "kind": db.EXEC_CONTRACT, "amount": b * 0.9, "doc_date": "2026-10-20"},
            {"id": 3, "site_code": fx.SITE, "item_code": self.lap["item_code"], "kind": db.EXEC_PAYMENT, "amount": b * 0.5, "doc_date": "2026-11-10"},
            {"id": 4, "site_code": fx.SITE, "item_code": self.zw["item_code"], "kind": db.EXEC_CONTRACT, "amount": self.zw["total_budget"] + 1e6, "doc_date": "2026-12-01"},
            {"id": 5, "site_code": fx.SITE, "item_code": "A01A27BDA04-05-999", "kind": db.EXEC_PAYMENT, "amount": 5e6, "doc_date": "2027-10-01"},
        ]

    def test_summary_flags_and_totals(self):
        s = ex.summarize(self.lines, self.E)
        r = s[s["item_code"] == self.lap["item_code"]].iloc[0]
        self.assertEqual((r["remaining"], r["flag"]), (0, ""))
        self.assertEqual(s[s["item_code"] == self.zw["item_code"]]["flag"].iloc[0], ex.FLAG_OVER)
        self.assertEqual(s[s["item_code"] == "A01A27BDA04-05-999"]["flag"].iloc[0], ex.FLAG_NO_LINE)
        self.assertNotEqual(s["flag"].iloc[0], "")  # cảnh báo lên đầu
        t = ex.totals(s)
        self.assertAlmostEqual(t["budget"], self.lines["total_budget"].sum())
        self.assertEqual(t["n_flag"], 2)

    def test_plan_vs_actual_and_checks(self):
        p = ex.plan_vs_actual(self.lines, self.E, fx.MONTHS)
        self.assertAlmostEqual(p["Kế hoạch"].sum(), self.lines["total_budget"].sum())
        self.assertEqual(p.loc[p["Tháng"] == "T11 2026", "Thực chi"].iloc[0], self.lap["total_budget"] * 0.5)
        self.assertEqual(p.attrs["outside"], 5e6)
        row = ex.summarize(self.lines, self.E).set_index("item_code").loc[self.lap["item_code"]].to_dict()
        self.assertEqual(ex.check_new(row, db.EXEC_CONTRACT, self.lap["total_budget"] * 0.1), "")
        self.assertIn("vượt ngân sách", ex.check_new(row, db.EXEC_CONTRACT, self.lap["total_budget"] * 0.2))
        self.assertIn("vượt hợp đồng", ex.check_new(row, db.EXEC_PAYMENT, self.lap["total_budget"] * 0.5))


class VersionDiffTest(unittest.TestCase):
    def test_added_removed_changed(self):
        a = fx.make_lines()
        b = copy.deepcopy(a)
        b[0]["quantity"] += 2
        b[0]["total_budget"] = b[0]["quantity"] * b[0]["unit_price"]
        gone = b.pop(1)
        b.append({"item_code": "A01A27BDA04-05-900", "item_name": "Máy mới", "dept_proposing": fx.KD, "quantity": 1,
                  "unit_price": 5e6, "total_budget": 5e6})
        d = vs.diff(a, b)
        st = dict(zip(d["item_code"], d["status"]))
        self.assertEqual((st[a[0]["item_code"]], st[gone["item_code"]], st["A01A27BDA04-05-900"]), (vs.CHANGED, vs.REMOVED, vs.ADDED))
        self.assertEqual(len(d), 3)
        self.assertIn("SL", d[d["item_code"] == a[0]["item_code"]]["what"].iloc[0])
        sm = vs.summary(d, a, b)
        self.assertAlmostEqual(sm["total_new"] - sm["total_old"], d["delta"].sum())
        self.assertAlmostEqual(vs.by_dept(a, b)["delta"].sum(), d["delta"].sum())
        self.assertTrue(vs.diff(a, a).empty)


class DatabaseTest(unittest.TestCase):
    """Các bảng mới trên CSDL SQLite tạm (tests/__init__.py đặt ICOST_DB_PATH)."""

    @classmethod
    def setUpClass(cls):
        db.init_db()
        db.replace_lines(fx.YEAR, fx.SITE, fx.make_lines(), actor="tester@icost.test")

    def test_dept_status_roundtrip(self):
        db.set_dept_status(fx.YEAR, fx.SITE, fx.KT, db.STATUS_SUBMITTED, "tester@icost.test")
        db.set_dept_status(fx.YEAR, fx.SITE, "  phòng kế toán ", db.STATUS_RETURNED, "tester@icost.test", "bổ sung căn cứ")
        st = db.dept_statuses(fx.YEAR, fx.SITE)
        self.assertEqual(len(st), 1)  # cùng phòng, không phân biệt hoa/thường
        self.assertEqual((st["phòng kế toán"]["status"], st["phòng kế toán"]["note"]), (db.STATUS_RETURNED, "bổ sung căn cứ"))
        self.assertEqual(db.get_dept_status(fx.YEAR, fx.SITE, fx.KD)["status"], db.STATUS_DRAFT)

    def test_exec_versions_and_backup_restore(self):
        code = db.load_lines(fx.YEAR, [fx.SITE])[0]["item_code"]
        db.add_exec(fx.YEAR, fx.SITE, code, db.EXEC_CONTRACT, 10e6, "tester@icost.test", "HD-1", "2026-10-20")
        with self.assertRaises(ValueError):
            db.add_exec(fx.YEAR, fx.SITE, code, "bogus", 1, "tester@icost.test")
        v1 = db.save_version(fx.YEAR, fx.SITE, db.load_lines(fx.YEAR, [fx.SITE]), "tester@icost.test")
        v2 = db.save_version(fx.YEAR, fx.SITE, db.load_lines(fx.YEAR, [fx.SITE]), "tester@icost.test", "điều chỉnh thử")
        self.assertEqual((v1["label"], v2["label"]), ("Bản duyệt", "Điều chỉnh lần 1"))
        self.assertEqual(db.latest_version_lines(fx.YEAR, [fx.SITE])[fx.SITE]["info"]["version_no"], v2["version_no"])
        backup = db.backup_bytes()
        db.delete_exec(db.load_exec(fx.YEAR, [fx.SITE])[0]["id"], "tester@icost.test")
        self.assertEqual(db.load_exec(fx.YEAR, [fx.SITE]), [])
        db.restore_bytes(backup)
        self.assertEqual(len(db.load_exec(fx.YEAR, [fx.SITE])), 1)
        self.assertEqual(len(db.list_versions(fx.YEAR, [fx.SITE])), 2)
        self.assertEqual(len(db.load_lines(fx.YEAR, [fx.SITE])), len(fx.SPEC))


if __name__ == "__main__":
    unittest.main()
