import io
import unittest
import pandas as pd
import openpyxl

from master_data import DEFAULT_MONTHS
from reports import purchasing_plan_df, purchasing_plan_summary, purchasing_plan_xlsx


class TestPurchasingPlan(unittest.TestCase):

    def setUp(self):
        self.months = DEFAULT_MONTHS  # 12 tháng: T10..T12 2025, T1..T9 2026

    def test_empty_dataframe(self):
        df_empty = pd.DataFrame()
        plan = purchasing_plan_df(df_empty, self.months)
        self.assertTrue(plan.empty)
        self.assertIn("qty_Q1", plan.columns)
        self.assertIn("amt_Q1", plan.columns)
        self.assertIn(f"qty_{self.months[0]}", plan.columns)

        summ = purchasing_plan_summary(df_empty, self.months)
        self.assertTrue(summ.empty)

        # Excel export should still return valid bytes
        xlsx = purchasing_plan_xlsx(df_empty, self.months, {"budget_year": "2026", "scope": "Test"})
        self.assertIsInstance(xlsx, bytes)
        self.assertGreater(len(xlsx), 1000)

        wb = openpyxl.load_workbook(io.BytesIO(xlsx), data_only=True)
        self.assertEqual(wb.sheetnames, ["Theo Tháng", "Theo Quý", "Tổng hợp Nhóm CNTT"])
        wb.close()

    def test_single_month_phasing(self):
        row = {
            "item_code": "NM-01-001",
            "item_name": "Laptop ThinkPad L14",
            "location": "AHA",
            "dept_proposing": "IT",
            "it_group": "IT01. Máy tính & Thiết bị văn phòng",
            "capex_type": "CAPEX",
            "unit": "Máy",
            "unit_price": 25_000_000,
            "quantity": 5,
            "total_budget": 125_000_000,
            f"pct_{self.months[0]}": 1.0,  # 100% vào T10 (Q1)
        }
        df = pd.DataFrame([row])
        plan = purchasing_plan_df(df, self.months)

        self.assertEqual(len(plan), 1)
        r0 = plan.iloc[0]
        self.assertEqual(r0[f"qty_{self.months[0]}"], 5)
        self.assertEqual(r0[f"amt_{self.months[0]}"], 125_000_000)
        self.assertEqual(r0["qty_Q1"], 5)
        self.assertEqual(r0["amt_Q1"], 125_000_000)
        self.assertEqual(r0["qty_Q2"], 0)
        self.assertEqual(r0["amt_Q2"], 0)
        self.assertEqual(r0["qty_Q3"], 0)
        self.assertEqual(r0["qty_Q4"], 0)

    def test_even_split_quarters_snap_units(self):
        # 10 máy chia đều 4 quý (T10, T1, T4, T7: tỷ lệ mỗi quý 25%)
        # snap_units(25% x 10) -> 3, 3, 2, 2
        pcts = {f"pct_{m}": 0.0 for m in self.months}
        pcts[f"pct_{self.months[0]}"] = 0.25  # T10 (Q1)
        pcts[f"pct_{self.months[3]}"] = 0.25  # T1 (Q2)
        pcts[f"pct_{self.months[6]}"] = 0.25  # T4 (Q3)
        pcts[f"pct_{self.months[9]}"] = 0.25  # T7 (Q4)

        row = {
            "item_code": "NM-01-002",
            "item_name": "Màn hình Dell 24 inch",
            "location": "LAN",
            "dept_proposing": "Kế toán",
            "it_group": "IT01. Máy tính & Thiết bị văn phòng",
            "capex_type": "CCDC",
            "unit": "Cái",
            "unit_price": 3_500_000,
            "quantity": 10,
            "total_budget": 35_000_000,
            **pcts,
        }
        df = pd.DataFrame([row])
        plan = purchasing_plan_df(df, self.months)

        r0 = plan.iloc[0]
        # Tổng SL các tháng phải bằng 10
        total_units = sum(r0[f"qty_{m}"] for m in self.months)
        self.assertEqual(total_units, 10)
        self.assertEqual(r0["qty_Q1"], 3)
        self.assertEqual(r0["qty_Q2"], 3)
        self.assertEqual(r0["qty_Q3"], 2)
        self.assertEqual(r0["qty_Q4"], 2)

        # Tổng tiền các tháng = 35_000_000. Tỷ lệ 25% không ứng SL nguyên (2,5 máy) -> tiền giữ theo dòng tiền ngân sách
        # (25% × 35tr = 8,75tr / quý, khớp biểu đồ giải ngân), SL làm tròn 3-3-2-2 chỉ để biết số máy cần mua
        total_amt = sum(r0[f"amt_{m}"] for m in self.months)
        self.assertEqual(total_amt, 35_000_000)
        for q in ("Q1", "Q2", "Q3", "Q4"):
            self.assertEqual(r0[f"amt_{q}"], 8_750_000)

        # Phân kỳ theo SL nguyên (cách app lưu khi chia đều SL: 30-30-20-20%) -> tiền = SL × đơn giá
        row.update({f"pct_{self.months[i]}": p for i, p in zip((0, 3, 6, 9), (0.3, 0.3, 0.2, 0.2))})
        r1 = purchasing_plan_df(pd.DataFrame([row]), self.months).iloc[0]
        self.assertEqual([r1[f"amt_{q}"] for q in ("Q1", "Q2", "Q3", "Q4")], [10_500_000, 10_500_000, 7_000_000, 7_000_000])

    def test_summary_and_xlsx_export(self):
        rows = [
            {
                "item_code": "IT-01",
                "item_name": "PC văn phòng",
                "location": "AHA",
                "dept_proposing": "Kinh doanh",
                "it_group": "IT01. Máy tính",
                "capex_type": "CAPEX",
                "unit": "Bộ",
                "unit_price": 20_000_000,
                "quantity": 4,
                "total_budget": 80_000_000,
                f"pct_{self.months[0]}": 1.0,
            },
            {
                "item_code": "IT-02",
                "item_name": "Switch Cisco Catalyst",
                "location": "LAN",
                "dept_proposing": "IT",
                "it_group": "IT03. Mạng & Hạ tầng",
                "capex_type": "CAPEX",
                "unit": "Cái",
                "unit_price": 45_000_000,
                "quantity": 2,
                "total_budget": 90_000_000,
                f"pct_{self.months[3]}": 0.5,
                f"pct_{self.months[6]}": 0.5,
            }
        ]
        df = pd.DataFrame(rows)

        summ = purchasing_plan_summary(df, self.months, "it_group")
        self.assertEqual(len(summ), 2)
        self.assertEqual(summ["Tổng ngân sách"].sum(), 170_000_000)
        self.assertEqual(summ["Tổng SL"].sum(), 6)

        xlsx_bytes = purchasing_plan_xlsx(df, self.months, {
            "budget_year": "2026",
            "scope": "Toàn tập đoàn",
            "date": "12/10/2026"
        })
        self.assertIsInstance(xlsx_bytes, bytes)
        self.assertGreater(len(xlsx_bytes), 2000)

        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        self.assertEqual(wb.sheetnames, ["Theo Tháng", "Theo Quý", "Tổng hợp Nhóm CNTT"])
        ws_m = wb["Theo Tháng"]
        self.assertEqual(ws_m.cell(row=6, column=2).value, "IT-01")
        self.assertEqual(ws_m.cell(row=7, column=2).value, "IT-02")
        self.assertEqual(ws_m.cell(row=8, column=1).value, "TỔNG CỘNG")

        ws_q = wb["Theo Quý"]
        self.assertEqual(ws_q.cell(row=8, column=1).value, "TỔNG CỘNG")

        ws_grp = wb["Tổng hợp Nhóm CNTT"]
        self.assertEqual(ws_grp.cell(row=7, column=1).value, "TỔNG CỘNG")
        wb.close()


class PlanAmountsMatchCashflowTest(unittest.TestCase):
    def test_uneven_ratio_keeps_budget_cashflow(self):
        """3 máy chia 50/50 (SL làm tròn 2-1): tiền tháng giữ theo ngân sách 15tr/15tr, không phải 20tr/10tr."""
        m = DEFAULT_MONTHS
        row = {"item_name": "Máy in", "quantity": 3, "unit_price": 10_000_000, "total_budget": 30_000_000,
               f"pct_{m[0]}": 0.5, f"pct_{m[1]}": 0.5, f"val_{m[0]}": 15_000_000, f"val_{m[1]}": 15_000_000}
        plan = purchasing_plan_df(pd.DataFrame([row]), m).iloc[0]
        self.assertEqual((plan[f"amt_{m[0]}"], plan[f"amt_{m[1]}"]), (15_000_000, 15_000_000))
        self.assertEqual(plan[f"qty_{m[0]}"] + plan[f"qty_{m[1]}"], 3)


if __name__ == "__main__":
    unittest.main()
