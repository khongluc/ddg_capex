"""
Báo cáo tổng hợp theo Khối: bảng Khối (so sánh năm trước) và file Excel Khối -> Phòng ban -> Hạng mục.
Khối lấy từ cột division của dòng ngân sách; trống -> theo bảng phòng ban -> khối (dept_divisions).
"""
import io
import datetime
from typing import Dict, Any, Optional, List

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from master_data import division_of, DEFAULT_MONTHS
from capex_engine import clean_number, snap_units, split_units, current_units

NO_DIVISION = "(Chưa có khối)"
NO_DEPT = "(Chưa ghi phòng ban)"
TYPES = ("CAPEX", "CCDC", "OPEX")
TYPE_LABEL = {"CAPEX": "TSCĐ (CAPEX)", "CCDC": "CCDC", "OPEX": "OPEX"}


def _key(v) -> str:
    return " ".join(str(v or "").upper().split())


def with_division(df: pd.DataFrame, master: Dict[str, Any]) -> pd.DataFrame:
    """Thêm cột khoi / dept / ctype đã chuẩn hóa (không đổi df gốc)."""
    if df is None or df.empty:
        return pd.DataFrame(columns=["khoi", "dept", "ctype", "total_budget"])
    d = df.copy()
    dept = d["dept_proposing"].fillna("").astype(str).str.strip() if "dept_proposing" in d.columns else pd.Series("", index=d.index)
    div = d["division"].fillna("").astype(str).str.strip() if "division" in d.columns else pd.Series("", index=d.index)
    div = div.where(~div.isin(["", "nan", "None"]), dept.map(lambda x: division_of(x, master)))
    d["khoi"] = div.map(lambda v: _key(v) or NO_DIVISION)
    d["dept"] = dept.where(dept != "", NO_DEPT)
    ct = d["capex_type"] if "capex_type" in d.columns else pd.Series("", index=d.index)
    d["ctype"] = ct.where(ct.isin(TYPES), "Chưa phân loại")
    d["total_budget"] = pd.to_numeric(d.get("total_budget", 0), errors="coerce").fillna(0.0)
    return d


def summary_by_division(df: pd.DataFrame, prev_df: Optional[pd.DataFrame], master: Dict[str, Any]) -> pd.DataFrame:
    """1 dòng / Khối: số phòng, số dòng, tiền theo CAPEX/CCDC/OPEX, tổng, tỷ trọng, năm trước, chênh lệch. Sắp theo tổng."""
    d = with_division(df, master)
    cols = ["Khối", "Số phòng", "Số dòng", *[TYPE_LABEL[t] for t in TYPES], "Chưa phân loại", "Tổng", "Tỷ trọng",
            "Năm trước", "Chênh lệch", "% so năm trước"]
    if d.empty:
        return pd.DataFrame(columns=cols)
    g = d.groupby("khoi")
    out = pd.DataFrame({"Số phòng": g["dept"].nunique(), "Số dòng": g.size(), "Tổng": g["total_budget"].sum()})
    pv = d.pivot_table(index="khoi", columns="ctype", values="total_budget", aggfunc="sum", fill_value=0.0)
    for t in (*TYPES, "Chưa phân loại"):
        out[TYPE_LABEL.get(t, t)] = pv[t] if t in pv.columns else 0.0
    total = out["Tổng"].sum()
    out["Tỷ trọng"] = out["Tổng"] / total if total else 0.0
    p = with_division(prev_df, master)
    prev = p.groupby("khoi")["total_budget"].sum() if not p.empty else pd.Series(dtype=float)
    out = out.reindex(out.index.union(prev.index), fill_value=0.0)
    out["Năm trước"] = prev.reindex(out.index).fillna(0.0)
    out["Chênh lệch"] = out["Tổng"] - out["Năm trước"]
    out["% so năm trước"] = (out["Chênh lệch"] / out["Năm trước"]).where(out["Năm trước"] > 0)
    out[["Số phòng", "Số dòng"]] = out[["Số phòng", "Số dòng"]].astype(int)
    out = out.sort_values(["Tổng", "Năm trước"], ascending=False)
    out.index.name = "Khối"
    return out.reset_index()[cols]


# ---------------------------------------------------------------- Excel
_THIN = Side(style="thin", color="CBD5E1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD = PatternFill("solid", fgColor="1E3A5F")
_FILL_K = PatternFill("solid", fgColor="DBEAFE")
_FILL_P = PatternFill("solid", fgColor="F1F5F9")
_MONEY = '#,##0'


def _header(ws, row: int, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), _HEAD, _BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _title(ws, title: str, sub: str):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color="1E3A5F")
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, color="64748B")


def report_xlsx(df: pd.DataFrame, prev_df: Optional[pd.DataFrame], master: Dict[str, Any], meta: Dict[str, Any]) -> bytes:
    year, scope = meta.get("budget_year", ""), meta.get("scope", "")
    prev_year = str(int(year) - 1) if str(year).isdigit() else "năm trước"
    sub = f"Phạm vi: {scope} · Đơn vị: VNĐ, chưa VAT, chưa gồm dự phòng · Lập ngày {meta.get('date', '')}"
    wb = Workbook()

    # Sheet 1: theo Khối
    ws = wb.active
    ws.title = "Theo Khối"
    _title(ws, f"TỔNG HỢP NGÂN SÁCH CNTT {year} THEO KHỐI", sub)
    s = summary_by_division(df, prev_df, master)
    titles = ["Khối", "Số phòng", "Số dòng", "TSCĐ (CAPEX)", "CCDC", "OPEX", "Chưa phân loại", f"Tổng {year}", "Tỷ trọng",
              f"Năm {prev_year}", "Chênh lệch", "% so năm trước"]
    _header(ws, 4, titles, [38, 10, 10, 18, 18, 18, 16, 20, 10, 18, 18, 12])
    keys = ["Khối", "Số phòng", "Số dòng", "TSCĐ (CAPEX)", "CCDC", "OPEX", "Chưa phân loại", "Tổng", "Tỷ trọng",
            "Năm trước", "Chênh lệch", "% so năm trước"]
    r = 5
    for rec in s.to_dict("records"):
        for i, k in enumerate(keys, 1):
            v = rec[k]
            c = ws.cell(row=r, column=i, value=None if (isinstance(v, float) and v != v) else v)
            c.border = _BORDER
            c.number_format = "0.0%" if k in ("Tỷ trọng", "% so năm trước") else (_MONEY if i > 1 else "@")
        r += 1
    if r > 5:  # dòng tổng bằng công thức
        ws.cell(row=r, column=1, value="TỔNG CỘNG")
        for i in range(2, len(keys) + 1):
            col = get_column_letter(i)
            if keys[i - 1] == "% so năm trước":
                val = f"=IF({get_column_letter(10)}{r}>0,{get_column_letter(11)}{r}/{get_column_letter(10)}{r},\"\")"
            elif keys[i - 1] == "Tỷ trọng":
                val = f"=SUM({col}5:{col}{r - 1})"
            else:
                val = f"=SUM({col}5:{col}{r - 1})"
            c = ws.cell(row=r, column=i, value=val)
            c.number_format = "0.0%" if keys[i - 1] in ("Tỷ trọng", "% so năm trước") else _MONEY
        for i in range(1, len(keys) + 1):
            ws.cell(row=r, column=i).font = Font(bold=True)
            ws.cell(row=r, column=i).fill = _FILL_K
            ws.cell(row=r, column=i).border = _BORDER
    ws.cell(row=r + 2, column=1, value="Số phòng ban của dòng TỔNG CỘNG là tổng theo từng khối.").font = Font(italic=True, color="64748B")

    # Sheet 2: Khối -> Phòng -> Hạng mục (gom theo mã danh mục + hình thức)
    ws2 = wb.create_sheet("Khối - Phòng - Hạng mục")
    _title(ws2, f"CHI TIẾT NGÂN SÁCH CNTT {year}: KHỐI → PHÒNG BAN → HẠNG MỤC", sub + " · Bấm dấu +/− bên trái để thu gọn")
    _header(ws2, 4, ["Khối / Phòng ban / Hạng mục", "Mã danh mục", "Hình thức", "Loại", "ĐVT", "Số lượng", "Đơn giá", "Thành tiền"],
            [56, 13, 18, 10, 10, 11, 16, 20])
    d = with_division(df, master)
    order = list(s["Khối"]) if not s.empty else []
    r = 5
    for khoi in order:
        dk = d[d["khoi"] == khoi]
        if dk.empty:
            continue
        rk = r
        ws2.cell(row=r, column=1, value=khoi)
        r += 1
        for dept, dp in sorted(dk.groupby("dept"), key=lambda x: -x[1]["total_budget"].sum()):
            rp = r
            ws2.cell(row=r, column=1, value=f"   {dept}")
            r += 1
            agg = (dp.assign(code=dp.get("catalog_code", pd.Series("", index=dp.index)).fillna(""),
                             inv=dp.get("invest_type", pd.Series("", index=dp.index)).fillna(""),
                             unit=dp.get("unit", pd.Series("", index=dp.index)).fillna(""),
                             price=pd.to_numeric(dp.get("unit_price", 0), errors="coerce").fillna(0.0),
                             qty=pd.to_numeric(dp.get("quantity", 0), errors="coerce").fillna(0.0))
                   .groupby(["item_name", "code", "inv", "ctype", "unit", "price"], dropna=False)
                   .agg(qty=("qty", "sum"), tot=("total_budget", "sum")).reset_index()
                   .sort_values("tot", ascending=False))
            first = r
            for it in agg.to_dict("records"):
                vals = [f"      {it['item_name']}", it["code"], it["inv"], it["ctype"], it["unit"], it["qty"], it["price"], it["tot"]]
                for i, v in enumerate(vals, 1):
                    c = ws2.cell(row=r, column=i, value=v)
                    c.border = _BORDER
                    if i >= 6:
                        c.number_format = _MONEY
                ws2.row_dimensions[r].outline_level = 2
                r += 1
            ws2.cell(row=rp, column=8, value=f"=SUM(H{first}:H{r - 1})")
            ws2.row_dimensions[rp].outline_level = 1
            for i in range(1, 9):
                c = ws2.cell(row=rp, column=i)
                c.font, c.fill, c.border = Font(bold=True), _FILL_P, _BORDER
                if i == 8:
                    c.number_format = _MONEY
        # tổng khối = tổng các dòng phòng ban của khối (cấp outline 1)
        dept_rows = [rr for rr in range(rk + 1, r) if ws2.row_dimensions[rr].outline_level == 1]
        ws2.cell(row=rk, column=8, value="=" + "+".join(f"H{rr}" for rr in dept_rows) if dept_rows else 0)
        for i in range(1, 9):
            c = ws2.cell(row=rk, column=i)
            c.font, c.fill, c.border = Font(bold=True, color="1E3A5F"), _FILL_K, _BORDER
            if i == 8:
                c.number_format = _MONEY
    if r > 5:
        khoi_rows = [rr for rr in range(5, r) if ws2.row_dimensions[rr].outline_level == 0]
        ws2.cell(row=r, column=1, value="TỔNG CỘNG")
        ws2.cell(row=r, column=8, value="=" + "+".join(f"H{rr}" for rr in khoi_rows))
        for i in range(1, 9):
            c = ws2.cell(row=r, column=i)
            c.font, c.fill, c.border = Font(bold=True), _FILL_K, _BORDER
            if i == 8:
                c.number_format = _MONEY
    ws2.sheet_properties.outlinePr.summaryBelow = False

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- Kế hoạch Mua sắm theo Tháng / Quý

def purchasing_plan_df(df: Optional[pd.DataFrame], months: Optional[List[str]] = None) -> pd.DataFrame:
    """
    Sinh bảng kế hoạch mua sắm và giải ngân theo 12 tháng và 4 quý cho từng hạng mục.
    Số lượng theo tháng được tính từ quantity × pct_<tháng> sử dụng snap_units để đảm bảo số nguyên nếu có thể.
    Tiền theo tháng tính từ val_<tháng> hoặc total_budget × pct_<tháng>.
    """
    if months is None:
        months = DEFAULT_MONTHS

    base_cols = [
        "item_code", "item_name", "location", "dept_proposing",
        "it_group", "capex_type", "unit", "unit_price", "quantity", "total_budget"
    ]
    quarter_qty_cols = ["qty_Q1", "qty_Q2", "qty_Q3", "qty_Q4"]
    quarter_amt_cols = ["amt_Q1", "amt_Q2", "amt_Q3", "amt_Q4"]
    month_qty_cols = [f"qty_{m}" for m in months]
    month_amt_cols = [f"amt_{m}" for m in months]
    all_cols = base_cols + quarter_qty_cols + quarter_amt_cols + month_qty_cols + month_amt_cols

    if df is None or df.empty:
        return pd.DataFrame(columns=all_cols)

    rows_out = []
    for _, r in df.iterrows():
        qty = clean_number(r.get("quantity", 0), 0.0)
        price = clean_number(r.get("unit_price", 0), 0.0)
        tot = clean_number(r.get("total_budget", 0), 0.0)
        if tot == 0.0 and qty > 0 and price > 0:
            tot = qty * price

        pcts = [clean_number(r.get(f"pct_{m}", 0.0), 0.0) for m in months]
        sum_pct = sum(pcts)
        # Nếu chưa phân kỳ nhưng có số lượng -> dồn vào tháng đầu tiên
        if sum_pct == 0.0 and qty > 0:
            pcts = [1.0] + [0.0] * (len(months) - 1)

        # Tính số lượng từng tháng
        is_int_qty = abs(qty - round(qty)) < 1e-9 and qty > 0
        if is_int_qty:
            units = current_units(pcts, qty)
        else:
            units = snap_units(pcts, qty) or split_units(qty, pcts)
            if units is None:
                if qty > 0 and sum(pcts) > 0:
                    units = [round(p * qty, 2) for p in pcts]
                else:
                    units = [0.0] * len(months)

        # Tính tiền từng tháng: SL nguyên khớp đúng phân kỳ -> tiền = SL × đơn giá; phân kỳ lẻ (vd. 50/50 cho 3 máy, SL làm tròn 2-1)
        # -> giữ tiền tháng của ngân sách (val_ / tỷ lệ) để khớp dòng tiền giải ngân ở các báo cáo khác
        exact_units = is_int_qty and snap_units(pcts, qty) is not None
        amts = []
        for i, m in enumerate(months):
            if exact_units and price > 0:
                amts.append(units[i] * price)
            else:
                val_col = f"val_{m}"
                if val_col in r and clean_number(r.get(val_col), 0.0) > 0:
                    amts.append(clean_number(r.get(val_col), 0.0))
                else:
                    amts.append(round(pcts[i] * tot, 0))

        # Gom 4 quý: Q1(0..2), Q2(3..5), Q3(6..8), Q4(9..11)
        q_qty = [
            sum(units[0:3]),
            sum(units[3:6]),
            sum(units[6:9]),
            sum(units[9:12])
        ]
        q_amt = [
            sum(amts[0:3]),
            sum(amts[3:6]),
            sum(amts[6:9]),
            sum(amts[9:12])
        ]

        row_dict = {
            "item_code": str(r.get("item_code") or "").strip(),
            "item_name": str(r.get("item_name") or "").strip(),
            "location": str(r.get("location") or "").strip(),
            "dept_proposing": str(r.get("dept_proposing") or "").strip(),
            "it_group": str(r.get("it_group") or "").strip(),
            "capex_type": str(r.get("capex_type") or "").strip(),
            "unit": str(r.get("unit") or "").strip(),
            "unit_price": price,
            "quantity": qty,
            "total_budget": tot,
            "qty_Q1": q_qty[0], "qty_Q2": q_qty[1], "qty_Q3": q_qty[2], "qty_Q4": q_qty[3],
            "amt_Q1": q_amt[0], "amt_Q2": q_amt[1], "amt_Q3": q_amt[2], "amt_Q4": q_amt[3],
        }
        for i, m in enumerate(months):
            row_dict[f"qty_{m}"] = units[i]
            row_dict[f"amt_{m}"] = amts[i]
        rows_out.append(row_dict)

    return pd.DataFrame(rows_out, columns=all_cols)


def purchasing_plan_summary(
    df: Optional[pd.DataFrame],
    months: Optional[List[str]] = None,
    group_by: str = "it_group"
) -> pd.DataFrame:
    """
    Bảng tổng hợp Kế hoạch mua sắm theo nhóm (Nhóm CNTT, Vị trí/Site, hoặc Phòng ban).
    """
    if months is None:
        months = DEFAULT_MONTHS
    plan = purchasing_plan_df(df, months)
    cols = [
        group_by, "Số hạng mục", "Tổng SL", "Tổng ngân sách",
        "SL Q1", "Tiền Q1", "SL Q2", "Tiền Q2", "SL Q3", "Tiền Q3", "SL Q4", "Tiền Q4"
    ]
    if plan.empty:
        return pd.DataFrame(columns=cols)

    grp = plan.groupby(group_by, dropna=False)
    summary = pd.DataFrame({
        "Số hạng mục": grp.size(),
        "Tổng SL": grp["quantity"].sum(),
        "Tổng ngân sách": grp["total_budget"].sum(),
        "SL Q1": grp["qty_Q1"].sum(),
        "Tiền Q1": grp["amt_Q1"].sum(),
        "SL Q2": grp["qty_Q2"].sum(),
        "Tiền Q2": grp["amt_Q2"].sum(),
        "SL Q3": grp["qty_Q3"].sum(),
        "Tiền Q3": grp["amt_Q3"].sum(),
        "SL Q4": grp["qty_Q4"].sum(),
        "Tiền Q4": grp["amt_Q4"].sum(),
    }).reset_index()

    summary = summary.sort_values("Tổng ngân sách", ascending=False)
    return summary[cols]


def purchasing_plan_xlsx(
    df: Optional[pd.DataFrame],
    months: Optional[List[str]] = None,
    meta: Optional[Dict[str, Any]] = None
) -> bytes:
    """
    Tạo file Excel kế hoạch mua sắm & giải ngân (3 sheet: Theo Tháng, Theo Quý, Tổng hợp Nhóm CNTT).
    """
    if months is None:
        months = DEFAULT_MONTHS
    meta = meta or {}
    year = meta.get("budget_year", "2026")
    scope = meta.get("scope", "Toàn bộ")
    today_str = meta.get("date", datetime.date.today().strftime("%d/%m/%Y"))
    sub_title = f"Phạm vi: {scope} · Đơn vị: VNĐ, chưa VAT · Lập ngày {today_str}"

    plan = purchasing_plan_df(df, months)
    wb = Workbook()

    # --- Sheet 1: Theo Tháng ---
    ws1 = wb.active
    ws1.title = "Theo Tháng"
    _title(ws1, f"KẾ HOẠCH MUA SẮM & GIẢI NGÂN CNTT THEO THÁNG NĂM {year}", sub_title)

    base_headers = ["STT", "Mã hạng mục", "Tên hạng mục", "Site", "Phòng ban", "Nhóm CNTT", "Phân loại", "ĐVT", "Đơn giá", "Tổng SL", "Thành tiền"]
    base_widths = [6, 15, 36, 12, 22, 22, 12, 10, 16, 12, 20]

    # Row 4 & 5: Header
    ws1.row_dimensions[4].height = 24
    ws1.row_dimensions[5].height = 24

    for c_idx, (th, w) in enumerate(zip(base_headers, base_widths), 1):
        for r_h in (4, 5):
            c = ws1.cell(row=r_h, column=c_idx)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = _HEAD
            c.border = _BORDER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws1.merge_cells(start_row=4, start_column=c_idx, end_row=5, end_column=c_idx)
        ws1.cell(row=4, column=c_idx, value=th)
        ws1.column_dimensions[get_column_letter(c_idx)].width = w

    for i, m in enumerate(months):
        start_c = 12 + i * 2
        for r_h in (4, 5):
            for c_offset in (0, 1):
                c = ws1.cell(row=r_h, column=start_c + c_offset)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = _HEAD
                c.border = _BORDER
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws1.merge_cells(start_row=4, start_column=start_c, end_row=4, end_column=start_c + 1)
        ws1.cell(row=4, column=start_c, value=m)
        ws1.cell(row=5, column=start_c, value="SL")
        ws1.cell(row=5, column=start_c + 1, value="Tiền")
        ws1.column_dimensions[get_column_letter(start_c)].width = 9
        ws1.column_dimensions[get_column_letter(start_c + 1)].width = 16

    # Fill data
    r = 6
    if not plan.empty:
        for idx, row in enumerate(plan.to_dict("records"), 1):
            ws1.cell(row=r, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws1.cell(row=r, column=2, value=row["item_code"])
            ws1.cell(row=r, column=3, value=row["item_name"])
            ws1.cell(row=r, column=4, value=row["location"])
            ws1.cell(row=r, column=5, value=row["dept_proposing"])
            ws1.cell(row=r, column=6, value=row["it_group"])
            ws1.cell(row=r, column=7, value=row["capex_type"]).alignment = Alignment(horizontal="center")
            ws1.cell(row=r, column=8, value=row["unit"]).alignment = Alignment(horizontal="center")
            c_p = ws1.cell(row=r, column=9, value=row["unit_price"])
            c_p.number_format = _MONEY
            c_q = ws1.cell(row=r, column=10, value=row["quantity"])
            c_q.number_format = '#,##0' if row["quantity"] == int(row["quantity"]) else '#,##0.0'
            c_t = ws1.cell(row=r, column=11, value=row["total_budget"])
            c_t.number_format = _MONEY

            for i, m in enumerate(months):
                start_c = 12 + i * 2
                u = row.get(f"qty_{m}", 0.0)
                amt = row.get(f"amt_{m}", 0.0)
                cu = ws1.cell(row=r, column=start_c, value=u if u != 0 else None)
                cu.number_format = '#,##0' if u == int(u) else '#,##0.0'
                ca = ws1.cell(row=r, column=start_c + 1, value=amt if amt != 0 else None)
                ca.number_format = _MONEY

            for col_i in range(1, 12 + len(months) * 2):
                ws1.cell(row=r, column=col_i).border = _BORDER
            r += 1

        # Dòng TỔNG CỘNG
        ws1.cell(row=r, column=1, value="TỔNG CỘNG")
        ws1.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws1.cell(row=r, column=10, value=f"=SUM(J6:J{r-1})").number_format = '#,##0'
        ws1.cell(row=r, column=11, value=f"=SUM(K6:K{r-1})").number_format = _MONEY
        for i in range(len(months)):
            start_c = 12 + i * 2
            col_u = get_column_letter(start_c)
            col_a = get_column_letter(start_c + 1)
            ws1.cell(row=r, column=start_c, value=f"=SUM({col_u}6:{col_u}{r-1})").number_format = '#,##0'
            ws1.cell(row=r, column=start_c + 1, value=f"=SUM({col_a}6:{col_a}{r-1})").number_format = _MONEY

        for col_i in range(1, 12 + len(months) * 2):
            cell = ws1.cell(row=r, column=col_i)
            cell.font = Font(bold=True)
            cell.fill = _FILL_K
            cell.border = _BORDER

    ws1.freeze_panes = "D6"

    # --- Sheet 2: Theo Quý ---
    ws2 = wb.create_sheet("Theo Quý")
    _title(ws2, f"KẾ HOẠCH MUA SẮM & GIẢI NGÂN CNTT THEO QUÝ NĂM {year}", sub_title)

    ws2.row_dimensions[4].height = 24
    ws2.row_dimensions[5].height = 24
    for c_idx, (th, w) in enumerate(zip(base_headers, base_widths), 1):
        for r_h in (4, 5):
            c = ws2.cell(row=r_h, column=c_idx)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = _HEAD
            c.border = _BORDER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws2.merge_cells(start_row=4, start_column=c_idx, end_row=5, end_column=c_idx)
        ws2.cell(row=4, column=c_idx, value=th)
        ws2.column_dimensions[get_column_letter(c_idx)].width = w

    q_labels = ["QUÝ 1 (T10-T12)", "QUÝ 2 (T1-T3)", "QUÝ 3 (T4-T6)", "QUÝ 4 (T7-T9)"]
    for i, q_lbl in enumerate(q_labels):
        start_c = 12 + i * 2
        for r_h in (4, 5):
            for c_offset in (0, 1):
                c = ws2.cell(row=r_h, column=start_c + c_offset)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = _HEAD
                c.border = _BORDER
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws2.merge_cells(start_row=4, start_column=start_c, end_row=4, end_column=start_c + 1)
        ws2.cell(row=4, column=start_c, value=q_lbl)
        ws2.cell(row=5, column=start_c, value="SL")
        ws2.cell(row=5, column=start_c + 1, value="Tiền")
        ws2.column_dimensions[get_column_letter(start_c)].width = 11
        ws2.column_dimensions[get_column_letter(start_c + 1)].width = 18

    r = 6
    if not plan.empty:
        for idx, row in enumerate(plan.to_dict("records"), 1):
            ws2.cell(row=r, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws2.cell(row=r, column=2, value=row["item_code"])
            ws2.cell(row=r, column=3, value=row["item_name"])
            ws2.cell(row=r, column=4, value=row["location"])
            ws2.cell(row=r, column=5, value=row["dept_proposing"])
            ws2.cell(row=r, column=6, value=row["it_group"])
            ws2.cell(row=r, column=7, value=row["capex_type"]).alignment = Alignment(horizontal="center")
            ws2.cell(row=r, column=8, value=row["unit"]).alignment = Alignment(horizontal="center")
            ws2.cell(row=r, column=9, value=row["unit_price"]).number_format = _MONEY
            ws2.cell(row=r, column=10, value=row["quantity"]).number_format = '#,##0' if row["quantity"] == int(row["quantity"]) else '#,##0.0'
            ws2.cell(row=r, column=11, value=row["total_budget"]).number_format = _MONEY

            for i, q_num in enumerate(["Q1", "Q2", "Q3", "Q4"]):
                start_c = 12 + i * 2
                u = row.get(f"qty_{q_num}", 0.0)
                amt = row.get(f"amt_{q_num}", 0.0)
                cu = ws2.cell(row=r, column=start_c, value=u if u != 0 else None)
                cu.number_format = '#,##0' if u == int(u) else '#,##0.0'
                ca = ws2.cell(row=r, column=start_c + 1, value=amt if amt != 0 else None)
                ca.number_format = _MONEY

            for col_i in range(1, 12 + 8):
                ws2.cell(row=r, column=col_i).border = _BORDER
            r += 1

        # Dòng TỔNG CỘNG
        ws2.cell(row=r, column=1, value="TỔNG CỘNG")
        ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws2.cell(row=r, column=10, value=f"=SUM(J6:J{r-1})").number_format = '#,##0'
        ws2.cell(row=r, column=11, value=f"=SUM(K6:K{r-1})").number_format = _MONEY
        for i in range(4):
            start_c = 12 + i * 2
            col_u = get_column_letter(start_c)
            col_a = get_column_letter(start_c + 1)
            ws2.cell(row=r, column=start_c, value=f"=SUM({col_u}6:{col_u}{r-1})").number_format = '#,##0'
            ws2.cell(row=r, column=start_c + 1, value=f"=SUM({col_a}6:{col_a}{r-1})").number_format = _MONEY

        for col_i in range(1, 12 + 8):
            cell = ws2.cell(row=r, column=col_i)
            cell.font = Font(bold=True)
            cell.fill = _FILL_K
            cell.border = _BORDER

    ws2.freeze_panes = "D6"

    # --- Sheet 3: Tổng hợp Nhóm CNTT ---
    ws3 = wb.create_sheet("Tổng hợp Nhóm CNTT")
    _title(ws3, f"TỔNG HỢP KẾ HOẠCH MUA SẮM THEO NHÓM CNTT NĂM {year}", sub_title)

    s3_headers = ["Nhóm CNTT", "Số hạng mục", "Tổng SL", "Tổng ngân sách",
                  "SL Q1", "Tiền Q1", "SL Q2", "Tiền Q2", "SL Q3", "Tiền Q3", "SL Q4", "Tiền Q4"]
    s3_widths = [30, 12, 12, 20, 11, 18, 11, 18, 11, 18, 11, 18]
    _header(ws3, 4, s3_headers, s3_widths)

    summ = purchasing_plan_summary(df, months, "it_group")
    r = 5
    if not summ.empty:
        for row in summ.to_dict("records"):
            ws3.cell(row=r, column=1, value=row.get("it_group", "(Chưa phân nhóm)"))
            ws3.cell(row=r, column=2, value=row.get("Số hạng mục", 0)).alignment = Alignment(horizontal="center")
            ws3.cell(row=r, column=3, value=row.get("Tổng SL", 0)).number_format = '#,##0'
            ws3.cell(row=r, column=4, value=row.get("Tổng ngân sách", 0)).number_format = _MONEY
            ws3.cell(row=r, column=5, value=row.get("SL Q1", 0)).number_format = '#,##0'
            ws3.cell(row=r, column=6, value=row.get("Tiền Q1", 0)).number_format = _MONEY
            ws3.cell(row=r, column=7, value=row.get("SL Q2", 0)).number_format = '#,##0'
            ws3.cell(row=r, column=8, value=row.get("Tiền Q2", 0)).number_format = _MONEY
            ws3.cell(row=r, column=9, value=row.get("SL Q3", 0)).number_format = '#,##0'
            ws3.cell(row=r, column=10, value=row.get("Tiền Q3", 0)).number_format = _MONEY
            ws3.cell(row=r, column=11, value=row.get("SL Q4", 0)).number_format = '#,##0'
            ws3.cell(row=r, column=12, value=row.get("Tiền Q4", 0)).number_format = _MONEY
            for col_i in range(1, 13):
                ws3.cell(row=r, column=col_i).border = _BORDER
            r += 1

        # Dòng TỔNG CỘNG
        ws3.cell(row=r, column=1, value="TỔNG CỘNG")
        ws3.cell(row=r, column=2, value=f"=SUM(B5:B{r-1})").alignment = Alignment(horizontal="center")
        for col_i in range(3, 13):
            col_letter = get_column_letter(col_i)
            c = ws3.cell(row=r, column=col_i, value=f"=SUM({col_letter}5:{col_letter}{r-1})")
            c.number_format = _MONEY if col_i in (4, 6, 8, 10, 12) else '#,##0'

        for col_i in range(1, 13):
            cell = ws3.cell(row=r, column=col_i)
            cell.font = Font(bold=True)
            cell.fill = _FILL_K
            cell.border = _BORDER

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

